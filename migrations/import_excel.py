"""
Data Migration — Import all data from MIS_Tracker_Data.xlsx into SQLite.
Run: python -m migrations.import_excel
"""
import os, sys, json
from datetime import datetime, date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app, db
from app.models import (
    User, AppSetting, ModuleToggle, Report, ReportComponent,
    CommodityConfig, TaskTemplate, MasterLog, KanbanTask, Sprint,
    RosterConfig, RosterLog, AuditLog
)

EXCEL_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                          'Old Data File', 'MIS_Tracker_Data.xlsx')


def parse_date(val):
    if val is None: return None
    if isinstance(val, datetime): return val.date()
    if isinstance(val, date): return val
    try: return datetime.fromisoformat(str(val)[:10]).date()
    except: return None


def parse_datetime(val):
    if val is None: return None
    if isinstance(val, datetime): return val
    try: return datetime.fromisoformat(str(val))
    except:
        try: return datetime.strptime(str(val), '%Y-%m-%d %H:%M:%S')
        except: return None


def run():
    import openpyxl
    print(f"Loading {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH)

    app = create_app()
    with app.app_context():
        db.drop_all()
        db.create_all()
        print("Database tables created.")

        # --- Users ---
        sh = wb['Users']
        count = 0
        for row in sh.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue
            email_val = str(row[0]).strip().lower()
            role_val = str(row[2] or 'Teammate')
            if email_val == 'nishant@agalabs.com':
                role_val = 'SuperAdmin'
            user = User(
                email=email_val,
                name=str(row[1] or ''),
                role=role_val,
                department=str(row[3] or 'Research'),
                active=str(row[4] or 'YES').upper() != 'NO',
                password_hash=str(row[5] or 'changeme'),
                last_login=parse_datetime(row[6]) if len(row) > 6 else None,
            )
            db.session.add(user)
            count += 1
        db.session.commit()
        print(f"  Users: {count}")

        # --- Config: Reports (col F-I, starting row 3) ---
        sh = wb['Config']
        count = 0
        for row in sh.iter_rows(min_row=3, min_col=6, max_col=9, values_only=True):
            if not row[0]: continue
            rpt = Report(
                name=str(row[0]),
                deadline=str(row[1] or ''),
                primary_assignee=str(row[2] or ''),
                frequency=str(row[3] or 'Daily'),
            )
            db.session.add(rpt)
            count += 1
        db.session.commit()
        print(f"  Reports: {count}")

        # --- Config: Components (col J-M, starting row 3) ---
        count = 0
        for row in sh.iter_rows(min_row=3, min_col=10, max_col=13, values_only=True):
            if not row[0] or not row[1]: continue
            comp = ReportComponent(
                report_name=str(row[0]),
                component_name=str(row[1]),
                assignee=str(row[2] or ''),
                min_count=int(row[3] or 0) if row[3] else 0,
            )
            db.session.add(comp)
            count += 1
        db.session.commit()
        print(f"  Report Components: {count}")

        # --- Config: Commodities (col P-T, starting row 3) ---
        count = 0
        for row in sh.iter_rows(min_row=3, min_col=16, max_col=20, values_only=True):
            if not row[0]: continue
            comm = CommodityConfig(
                commodity=str(row[0]),
                type=str(row[1] or ''),
                factor=str(row[2] or ''),
                mode=str(row[3] or ''),
                source=str(row[4] or ''),
            )
            db.session.add(comm)
            count += 1
        db.session.commit()
        print(f"  Commodity Configs: {count}")

        # --- Config: Tasks (col V-Z, starting row 3) ---
        count = 0
        for row in sh.iter_rows(min_row=3, min_col=22, max_col=26, values_only=True):
            if not row[0]: continue
            task = TaskTemplate(
                name=str(row[0]),
                assignee=str(row[1] or ''),
                recurrence=str(row[3] or 'daily').lower(),
                due_date=parse_datetime(row[4]) if row[4] else None,
            )
            db.session.add(task)
            count += 1
        db.session.commit()
        print(f"  Task Templates: {count}")

        # --- Config: Settings (col AB-AC, starting row 3) ---
        count = 0
        for row in sh.iter_rows(min_row=3, min_col=28, max_col=29, values_only=True):
            if not row[0]: continue
            setting = AppSetting(key=str(row[0]).strip(), value=str(row[1] or ''))
            db.session.add(setting)
            count += 1
        db.session.commit()
        print(f"  Settings: {count}")

        # --- Config: Modules (col AH-AI, starting row 3) ---
        count = 0
        for row in sh.iter_rows(min_row=3, min_col=34, max_col=35, values_only=True):
            if not row[0]: continue
            toggle = ModuleToggle(
                module_key=str(row[0]).strip().lower(),
                enabled=str(row[1] or 'YES').upper() != 'NO',
            )
            db.session.add(toggle)
            count += 1
        db.session.commit()
        print(f"  Module Toggles: {count}")

        # --- Master_Log ---
        sh = wb['Master_Log']
        count = 0
        for row in sh.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue
            entry = MasterLog(
                log_id=str(row[0]),
                timestamp_submitted=parse_datetime(row[1]) or datetime.now(),
                submitted_by_email=str(row[2] or '').lower(),
                on_behalf_of_email=str(row[3] or '').lower(),
                on_behalf_of_name=str(row[4] or ''),
                entry_type=str(row[5] or ''),
                entry_date=parse_date(row[6]) or date.today(),
                week_number=int(row[7] or 0) if row[7] else 0,
                month_year=str(row[8] or ''),
                topic=str(row[9] or ''),
                sub_topic=str(row[10] or ''),
                status=str(row[11] or 'Pending'),
                data_count=int(row[12] or 0) if row[12] else 0,
                mode_or_extra=str(row[13] or ''),
                remarks=str(row[14] or ''),
                is_late=str(row[15] or 'No') == 'Yes',
                time_to_submit_min=int(row[16] or 0) if row[16] else 0,
                row_status=str(row[17] or 'Active'),
                source_url=str(row[18] or '') if len(row) > 18 else '',
                confidence=str(row[19] or 'Medium') if len(row) > 19 else 'Medium',
                validated_by_email=str(row[20] or '') if len(row) > 20 else '',
            )
            db.session.add(entry)
            count += 1
        db.session.commit()
        print(f"  Master Log: {count}")

        # --- Kanban Tasks ---
        sh = wb['Kanban Tasks']
        count = 0
        for row in sh.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue
            task = KanbanTask(
                task_id=str(row[0]),
                title=str(row[1] or ''),
                description=str(row[2] or ''),
                assigned_to=str(row[3] or '').lower(),
                assigned_to_name=str(row[4] or ''),
                status=str(row[5] or 'Assigned'),
                due_date=parse_date(row[7]),
                created_by=str(row[8] or '').lower(),
                tags=str(row[9] or '') if len(row) > 9 else '',
                sprint_id=str(row[10] or '') if len(row) > 10 else '',
                story_points=int(row[11] or 0) if len(row) > 11 and row[11] else 0,
                department=str(row[12] or '') if len(row) > 12 else '',
                remark=str(row[19] or '') if len(row) > 19 else '',
            )
            db.session.add(task)
            count += 1
        db.session.commit()
        print(f"  Kanban Tasks: {count}")

        # --- Sprints ---
        sh = wb['Kanban Sprints']
        count = 0
        for row in sh.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue
            sprint = Sprint(
                sprint_id=str(row[0]),
                name=str(row[1] or ''),
                start_date=parse_date(row[2]),
                end_date=parse_date(row[3]),
                status=str(row[4] or 'Active'),
                goal=str(row[5] or ''),
            )
            db.session.add(sprint)
            count += 1
        db.session.commit()
        print(f"  Sprints: {count}")

        # --- Roster ---
        sh = wb['Daily_Roadster']
        count = 0
        for row in sh.iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[1]: continue
            rc = RosterConfig(
                role=str(row[0]).strip(),
                employee=str(row[1]).strip(),
                frequency=str(row[2] or 'Daily').strip(),
                remark=str(row[3] or '').strip(),
            )
            db.session.add(rc)
            count += 1
        db.session.commit()
        print(f"  Roster Configs: {count}")

        # --- Roster Log ---
        sh = wb['Roadster_Log']
        count = 0
        for row in sh.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue
            rl = RosterLog(
                date=parse_date(row[0]) or date.today(),
                role=str(row[1] or ''),
                employee=str(row[2] or ''),
                status=str(row[3] or 'Pending'),
                updated_at=parse_datetime(row[4]),
            )
            db.session.add(rl)
            count += 1
        db.session.commit()
        print(f"  Roster Log: {count}")

        # --- Audit Log ---
        sh = wb['Audit_Log']
        count = 0
        for row in sh.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue
            audit = AuditLog(
                audit_id=str(row[0]),
                timestamp=parse_datetime(row[1]) or datetime.now(),
                actor_email=str(row[2] or ''),
                action=str(row[3] or ''),
                target_type=str(row[4] or ''),
                target_id=str(row[5] or ''),
                before_json=str(row[6] or ''),
                after_json=str(row[7] or ''),
                is_proxy=bool(row[8]) if len(row) > 8 and row[8] else False,
            )
            db.session.add(audit)
            count += 1
        db.session.commit()
        print(f"  Audit Log: {count}")

        print("\n[OK] Migration complete!")
        print(f"   Database: {app.config['SQLALCHEMY_DATABASE_URI']}")


if __name__ == '__main__':
    run()
